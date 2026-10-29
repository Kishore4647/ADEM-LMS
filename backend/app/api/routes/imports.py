import csv
import io

from fastapi import APIRouter, Depends, HTTPException
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes.auth import get_current_user
from app.api.routes.workflows import STORAGE, get_record, write_mark
from app.core.database import get_db
from app.models import AssessmentComponent, MarkRecord, StudentProfile, User
from app.models.imports import ImportJob
from app.models.workflows import PrivateFile
from app.schemas.workflows import MarkInput, Payload
from app.services.access import audit, check_school, record_dict, require_role

router = APIRouter(prefix="/workspace/imports", tags=["external marks imports"])


class ImportInput(Payload):
    file_id: int


@router.get("")
def list_imports(user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    require_role(user, "super_admin", "content_uploader")
    return [
        record_dict(j)
        for j in db.scalars(
            select(ImportJob)
            .where(
                True
                if user.role.value == "super_admin"
                else ImportJob.owner_id == user.id
            )
            .order_by(ImportJob.id.desc())
        )
    ]


@router.post("/preview")
def preview(
    payload: ImportInput,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_role(user, "super_admin", "content_uploader")
    file = get_record(db, PrivateFile, payload.file_id)
    if file.owner_id != user.id:
        raise HTTPException(403, "File access denied.")
    data = (STORAGE / file.storage_key).read_bytes()
    if file.name.lower().endswith(".csv"):
        try:
            rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        except UnicodeDecodeError:
            raise HTTPException(400, "CSV must use UTF-8 encoding.")
    elif file.name.lower().endswith(".xlsx"):
        from zipfile import ZipFile

        from openpyxl import load_workbook

        with ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 40_000_000:
                raise HTTPException(400, "Expanded spreadsheet exceeds limit.")
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            raise HTTPException(400, "Spreadsheet is empty.")
        rows = [dict(zip(headers, row)) for row in iterator]
    else:
        raise HTTPException(400, "Use CSV or XLSX.")
    if not rows or len(rows) > 2000:
        raise HTTPException(400, "Import between 1 and 2000 rows.")
    parsed = []
    seen = set()
    for i, row in enumerate(rows, 2):
        try:
            email = str(row.get("student_email", "")).lower().strip()
            account = db.scalar(select(User).where(User.email == email))
            student = (
                db.scalar(
                    select(StudentProfile).where(StudentProfile.user_id == account.id)
                )
                if account
                else None
            )
            if not student:
                raise ValueError("Unmatched student_email.")
            check_school(db, user, student.school_id)
            values = dict(
                student_id=student.id,
                subject_node_id=int(row["subject_id"]),
                component=str(row["component"]),
                academic_period=str(row["academic_period"]),
                title=str(row["title"]),
                obtained_marks=None
                if str(row.get("absent", "")).lower() in {"true", "yes", "1"}
                else float(row["marks"]),
                maximum_marks=float(row["maximum"]),
                is_absent=str(row.get("absent", "")).lower() in {"true", "yes", "1"},
            )
            mark = MarkInput(**values)
            from app.services.access import subject_class

            if subject_class(db, mark.subject_node_id) != student.class_node_id:
                raise ValueError("Subject does not belong to student class.")
            key = (
                student.id,
                mark.subject_node_id,
                mark.academic_period,
                mark.component,
                "exam" if mark.component == "quarterly_exam" else mark.title,
            )
            if key in seen:
                raise ValueError("Duplicate row in uploaded file.")
            seen.add(key)
            existing = db.scalar(
                select(MarkRecord).where(
                    MarkRecord.student_id == student.id,
                    MarkRecord.subject_node_id == mark.subject_node_id,
                    MarkRecord.academic_period == mark.academic_period,
                    MarkRecord.component == AssessmentComponent(mark.component),
                    True
                    if mark.component == "quarterly_exam"
                    else MarkRecord.title == mark.title,
                )
            )
            if existing:
                raise ValueError(
                    "Existing mark: use authorised mark correction instead of importing again."
                )
            parsed.append(
                dict(row=i, status="valid", values=mark.model_dump(), student=email)
            )
        except (ValueError, KeyError, TypeError, ValidationError, HTTPException) as e:
            parsed.append(
                dict(
                    row=i,
                    status="invalid",
                    error=str(e.detail) if isinstance(e, HTTPException) else str(e),
                )
            )
    job = ImportJob(
        owner_id=user.id,
        file_id=file.id,
        rows=parsed,
        summary=dict(
            valid=sum(r["status"] == "valid" for r in parsed),
            invalid=sum(r["status"] == "invalid" for r in parsed),
        ),
    )
    db.add(job)
    db.flush()
    audit(db, user, "validate", "marks_import", job.id, job.summary)
    db.commit()
    return record_dict(job)


@router.post("/{id}/commit")
def commit(
    id: int, user: User = Depends(get_current_user), db: Session = Depends(get_db)
):
    require_role(user, "super_admin", "content_uploader")
    job = db.scalar(select(ImportJob).where(ImportJob.id == id).with_for_update())
    if not job:
        raise HTTPException(404, "Import job not found.")
    if job.owner_id != user.id and user.role.value != "super_admin":
        raise HTTPException(403, "Import access denied.")
    if job.status != "validated":
        raise HTTPException(409, "Import already processed.")
    successful = 0
    rows = []
    for row in job.rows:
        if row["status"] != "valid":
            rows.append(row)
            continue
        try:
            with db.begin_nested():
                write_mark(db, user, MarkInput(**row["values"]))
            successful += 1
            rows.append({**row, "status": "imported"})
        except (HTTPException, ValueError) as e:
            rows.append(
                {
                    **row,
                    "status": "failed",
                    "error": str(e.detail) if isinstance(e, HTTPException) else str(e),
                }
            )
    job.rows = rows
    job.status = "completed"
    job.summary = dict(
        successful=successful,
        failed=sum(r["status"] == "failed" for r in rows),
        skipped=sum(r["status"] == "invalid" for r in rows),
    )
    audit(db, user, "commit", "marks_import", id, job.summary)
    db.commit()
    return record_dict(job)
